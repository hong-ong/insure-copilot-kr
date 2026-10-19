"""파일 형식별로 텍스트를 뽑는다 (PDF·HWP·HWPX·DOCX·TXT/MD).

여기서는 "글자를 꺼내는 일"만 한다. 꺼낸 글자가 제대로인지(한글 인식률)는 quality.py,
깨졌을 때 이미지로 다시 읽는 일은 ocr.py가 맡는다.
"""

from __future__ import annotations

import io
import re
import struct
import zipfile
import zlib
from dataclasses import dataclass, field
from xml.etree import ElementTree


class UnsupportedFormat(ValueError):
    """지원하지 않는 형식이거나 열 수 없는 파일."""


@dataclass
class Extracted:
    """추출 결과. pages는 페이지(또는 구역)별 텍스트로, 출처의 페이지 번호에 쓴다."""

    file_format: str
    pages: list[str]
    issues: list[str] = field(default_factory=list)
    is_scanned_like: bool = False

    @property
    def text(self) -> str:
        """TODO(재홍): 무엇을 왜 하는가 한 문장"""
        return "\n\n".join(self.pages)


EXT_FORMATS = {
    ".pdf": "pdf",
    ".hwp": "hwp",
    ".hwpx": "hwpx",
    ".docx": "docx",
    ".txt": "txt",
    ".md": "md",
}


def detect_format(file_name: str, data: bytes) -> str:
    """TODO(재홍): 무엇을 왜 하는가 한 문장"""
    # 확장자보다 파일 앞부분(매직 넘버)을 먼저 믿는다. 확장자는 사용자가 바꿀 수 있다.
    if data[:5] == b"%PDF-":
        return "pdf"
    if data[:8] == bytes.fromhex("D0CF11E0A1B11AE1"):
        return "hwp"  # OLE 복합 문서 (HWP 5.x)
    if data[:2] == b"PK":
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            names = z.namelist()
        if any(n.startswith("Contents/section") for n in names):
            return "hwpx"
        if "word/document.xml" in names:
            return "docx"
    ext = "." + file_name.rsplit(".", 1)[-1].lower() if "." in file_name else ""
    fmt = EXT_FORMATS.get(ext)
    if fmt in ("txt", "md"):
        return fmt
    raise UnsupportedFormat(f"지원하지 않는 형식입니다: {file_name}")


def extract(file_name: str, data: bytes) -> Extracted:
    """TODO(재홍): 무엇을 왜 하는가 한 문장"""
    fmt = detect_format(file_name, data)
    return {
        "pdf": _extract_pdf,
        "hwp": _extract_hwp,
        "hwpx": _extract_hwpx,
        "docx": _extract_docx,
        "txt": _extract_text,
        "md": _extract_text,
    }[fmt](data, fmt)


# ───────────────────────── PDF ─────────────────────────


def _extract_pdf(data: bytes, fmt: str) -> Extracted:
    """TODO(재홍): 무엇을 왜 하는가 한 문장"""
    import pymupdf

    with pymupdf.open(stream=data, filetype="pdf") as doc:
        if doc.needs_pass:
            raise UnsupportedFormat("암호가 걸린 PDF는 열 수 없습니다.")
        pages = [page.get_text("text") for page in doc]
        image_pages = sum(
            1 for page in doc if page.get_images() and len(page.get_text().strip()) < 30
        )
    result = Extracted(file_format=fmt, pages=pages)
    if pages and image_pages / len(pages) > 0.5:
        result.is_scanned_like = True
        result.issues.append(f"스캔 이미지로 보이는 페이지 {image_pages}/{len(pages)}")
    return result


def render_pdf_pages(data: bytes, max_pages: int, dpi: int = 150) -> list[bytes]:
    """TODO(재홍): 무엇을 왜 하는가 한 문장"""
    import pymupdf

    images = []
    with pymupdf.open(stream=data, filetype="pdf") as doc:
        for i, page in enumerate(doc):
            if i >= max_pages:
                break
            images.append(page.get_pixmap(dpi=dpi).tobytes("png"))
    return images


# ───────────────────────── HWP 5.x ─────────────────────────
# HWP 5.x = OLE 복합 문서. 본문은 BodyText/Section0, Section1 ... 스트림에 "레코드" 단위로 들어 있다.
# 레코드 헤더(4바이트) = tag(10비트) | level(10비트) | size(12비트). 문단 글자는 tag 67(HWPTAG_PARA_TEXT).

_HWPTAG_PARA_TEXT = 67
# 8글자(16바이트)를 차지하는 확장/인라인 컨트롤 문자
_HWP_WIDE_CONTROLS = {1, 2, 3, 4, 5, 6, 7, 8, 9, 11, 12, 14, 15, 16, 17, 18, 19, 20, 21, 22, 23}


def _extract_hwp(data: bytes, fmt: str) -> Extracted:
    """TODO(재홍): 무엇을 왜 하는가 한 문장"""
    import olefile

    ole = olefile.OleFileIO(io.BytesIO(data))
    try:
        header = ole.openstream("FileHeader").read()
        flags = struct.unpack("<I", header[36:40])[0]
        compressed = bool(flags & 0x01)
        if flags & 0x02:
            raise UnsupportedFormat("암호가 걸린 HWP는 열 수 없습니다.")
        if flags & 0x04:
            raise UnsupportedFormat("배포용 HWP(복사 방지)는 열 수 없습니다.")
        sections = sorted(
            (e for e in ole.listdir() if len(e) == 2 and e[0] == "BodyText"),
            key=lambda e: int(re.sub(r"\D", "", e[1]) or 0),
        )
        pages = []
        for entry in sections:
            raw = ole.openstream(entry).read()
            if compressed:
                raw = zlib.decompress(raw, -15)
            pages.append(_hwp_section_text(raw))
    finally:
        ole.close()
    if not pages:
        raise UnsupportedFormat("HWP 본문을 찾지 못했습니다.")
    return Extracted(file_format=fmt, pages=pages)


def _hwp_section_text(raw: bytes) -> str:
    """TODO(재홍): 무엇을 왜 하는가 한 문장"""
    out: list[str] = []
    pos = 0
    while pos + 4 <= len(raw):
        header = struct.unpack("<I", raw[pos : pos + 4])[0]
        tag, size = header & 0x3FF, (header >> 20) & 0xFFF
        pos += 4
        if size == 0xFFF:
            size = struct.unpack("<I", raw[pos : pos + 4])[0]
            pos += 4
        if tag == _HWPTAG_PARA_TEXT:
            out.append(_hwp_para_text(raw[pos : pos + size]))
        pos += size
    return "\n".join(line for line in out if line.strip())


def _hwp_para_text(body: bytes) -> str:
    """TODO(재홍): 무엇을 왜 하는가 한 문장"""
    chars: list[str] = []
    i = 0
    while i + 1 < len(body):
        code = struct.unpack("<H", body[i : i + 2])[0]
        if code in _HWP_WIDE_CONTROLS:
            i += 16
            continue
        if code < 32:
            if code in (10, 13):
                chars.append("\n")
            elif code == 9:
                chars.append("\t")
            i += 2
            continue
        chars.append(chr(code))
        i += 2
    return "".join(chars).strip()


# ───────────────────────── HWPX / DOCX ─────────────────────────


def _extract_hwpx(data: bytes, fmt: str) -> Extracted:
    """TODO(재홍): 무엇을 왜 하는가 한 문장"""
    pages = []
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        names = sorted(
            (n for n in z.namelist() if re.match(r"Contents/section\d+\.xml$", n)),
            key=lambda n: int(re.sub(r"\D", "", n)),
        )
        for name in names:
            root = ElementTree.fromstring(z.read(name))
            paragraphs = []
            for p in root.iter():
                if p.tag.endswith("}p"):
                    text = "".join(t.text or "" for t in p.iter() if t.tag.endswith("}t"))
                    if text.strip():
                        paragraphs.append(text.strip())
            pages.append("\n".join(dict.fromkeys(paragraphs)))  # 표 안 문단 중복 제거
    return Extracted(file_format=fmt, pages=pages)


def _extract_docx(data: bytes, fmt: str) -> Extracted:
    """TODO(재홍): 무엇을 왜 하는가 한 문장"""
    import docx

    document = docx.Document(io.BytesIO(data))
    lines = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return Extracted(file_format=fmt, pages=["\n".join(lines)])


def _extract_text(data: bytes, fmt: str) -> Extracted:
    """TODO(재홍): 무엇을 왜 하는가 한 문장"""
    for encoding in ("utf-8-sig", "cp949"):
        try:
            return Extracted(file_format=fmt, pages=[data.decode(encoding)])
        except UnicodeDecodeError:
            continue
    return Extracted(
        file_format=fmt,
        pages=[data.decode("utf-8", errors="replace")],
        issues=["문자 인코딩을 알 수 없어 일부 글자가 깨졌습니다."],
    )
